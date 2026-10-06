"""Build the Design A construction set: PDF, SVG sheets, DXF, BOQ.xlsx."""
import os
import ezdxf
from ezdxf import units
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
import build
from cad import DXF_LAYERS
import sheets_a, sheets_b, sheets_c, sheets_d

OUT = build.OUT


def new_dxf():
    doc = ezdxf.new("R2018", setup=True)
    doc.units = units.MM
    for name, col in DXF_LAYERS.items():
        if name not in doc.layers:
            doc.layers.add(name, color=col)
    return doc


def main():
    doc = new_dxf()
    msp = doc.modelspace()
    order = [
        ("A-000", "Cover sheet, drawing register & project information"),
        ("A-001", "General notes, assumptions & outline specification"),
        ("A-002", "Construction sequence, hold points & programme"),
        ("A-100", "Courtyard setting-out & general arrangement plan 1:50"),
        ("A-101", "Paving layout, levels & surface water drainage plan 1:50"),
        ("A-102", "Canopy roof plan — falls, gutter & membrane 1:25"),
        ("A-103", "Reflected ceiling plan, lighting & power 1:25 / 1:100"),
        ("A-104", "Planting & irrigation plan 1:25"),
        ("A-200", "Elevation A — annex façade & canopy 1:25"),
        ("A-201", "Elevations B & C 1:50"),
        ("A-300", "Section A-A — canopy, threshold & paving 1:20"),
        ("A-301", "Section B-B — planter, canopy & courtyard 1:25"),
        ("A-302", "Section C-C planter; slot drain, gully & edge details 1:10"),
        ("A-500", "Details 1 — post base, corner, eave & façade junction"),
        ("A-501", "Details 2 — LED, soffit, back-span, fascia & wires"),
        ("A-600", "Door & screen schedule, elevations & joinery"),
        ("S-001", "Structural design basis, loads & member checks"),
        ("S-100", "Canopy steel framing plan 1:25"),
        ("S-101", "Connections, foundation plan & member schedule"),
        ("Q-001", "Bill of quantities"),
    ]
    sheets = [
        sheets_d.sheet_a000(order), sheets_d.sheet_a001(), sheets_d.sheet_a002(),
        sheets_a.sheet_a100(msp), sheets_a.sheet_a101(msp), sheets_a.sheet_a102(msp), sheets_a.sheet_a103(msp),
        sheets_a.sheet_a104(),
        sheets_b.sheet_a200(msp), sheets_b.sheet_a201(msp), sheets_b.sheet_a300(msp), sheets_b.sheet_a301(msp),
        sheets_b.sheet_a302(msp),
        sheets_c.sheet_a500(msp), sheets_c.sheet_a501(), sheets_c.sheet_a600(),
        sheets_d.sheet_s001(), sheets_d.sheet_s100(msp), sheets_d.sheet_s101(), sheets_d.sheet_q001(),
    ]
    pdf = build.to_pdf(sheets, "Design-A_Construction-Set")
    build.to_svgs(sheets, "Design-A_svg")
    doc.saveas(os.path.join(OUT, "Design-A_CAD-1to1.dxf"))
    # BOQ
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "BOQ Design A"
    ws.append(["Item", "Description", "Unit", "Qty", "Rate", "Total"])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="333333")
    for r in sheets_d.boq_rows():
        ws.append([r[0], r[1], r[2], r[3], None, None])
        row = ws.max_row
        if r[0].isdigit() and "." not in r[0]:
            for c in ws[row]:
                c.font = Font(bold=True)
                c.fill = PatternFill("solid", fgColor="EFECE6")
        else:
            try:
                float(r[3])
                ws.cell(row=row, column=6, value=f"=IFERROR(D{row}*E{row},\"\")")
            except ValueError:
                pass
    ws.append([])
    ws.append(["", "SUBTOTAL", "", "", "", f"=SUM(F2:F{ws.max_row})"])
    ws.column_dimensions["A"].width = 7
    ws.column_dimensions["B"].width = 70
    ws.column_dimensions["C"].width = 7
    ws.column_dimensions["D"].width = 12
    ws.column_dimensions["E"].width = 12
    ws.column_dimensions["F"].width = 14
    wb.save(os.path.join(OUT, "BOQ_Design-A.xlsx"))
    print("PDF:", pdf)
    return sheets


if __name__ == "__main__":
    main()
